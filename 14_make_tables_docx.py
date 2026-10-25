import os
# -*- coding: utf-8 -*-
"""14_make_tables_docx.py —— 从 03_manuscript.md 提取三个表，生成 BJA 投稿用 Word 表格文件"""
import re
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = os.environ.get("MS_MD", "./03_manuscript.md")
OUT = os.environ.get("MS_DOCX", "./05_tables.docx")
md = open(SRC, encoding="utf-8").read()


def grab_table(n):
    """返回 (caption, rows, footnote)"""
    start = md.index(f"**Table {n}**")
    block = md[start:]
    # 到下一个 **Table 或 ### 为止
    end = len(block)
    for marker in ["**Table ", "\n### ", "\n## "]:
        k = block.find(marker, 10)
        if k != -1:
            end = min(end, k)
    block = block[:end]

    lines = block.strip().split("\n")
    caption = lines[0].replace(f"**Table {n}**", "").strip()
    rows, foot_parts = [], []
    for ln in lines[1:]:
        t = ln.strip()
        if t.startswith("|"):
            cells = [c.strip() for c in t.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                continue
            rows.append(cells)
        elif t:
            # 表注只取表后第一段；遇到第二段正文即停止，避免把结果段塞进表注
            if foot_parts:
                break
            if t.startswith("**") or t.startswith("#"):
                break
            foot_parts.append(t)
    footnote = " ".join(foot_parts).strip()
    return caption, rows, footnote


def set_cell_border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "bottom"):
        if edge in kw:
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(kw[edge]))
            el.set(qn("w:color"), "000000")
            borders.append(el)
    tcPr.append(borders)


doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(9)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2)
    s.left_margin = s.right_margin = Cm(2)

for n in (1, 2, 3):
    cap, rows, foot = grab_table(n)

    p = doc.add_paragraph()
    r = p.add_run(f"Table {n}  ")
    r.bold = True; r.font.size = Pt(9)
    r2 = p.add_run(cap)
    r2.font.size = Pt(9)
    p.paragraph_format.space_after = Pt(4)

    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = True

    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ""
            para = cell.paragraphs[0]
            para.paragraph_format.space_before = Pt(1)
            para.paragraph_format.space_after = Pt(1)
            run = para.add_run(val)
            run.font.size = Pt(8.5)
            run.font.name = "Times New Roman"
            if i == 0:
                run.bold = True
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                set_cell_border(cell, top=12, bottom=6)
            else:
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                if i == len(rows) - 1:
                    set_cell_border(cell, bottom=12)

    if foot:
        fp = doc.add_paragraph()
        fr = fp.add_run(foot)
        fr.font.size = Pt(7.5)
        fr.font.name = "Times New Roman"
        fp.paragraph_format.space_before = Pt(3)
        fp.paragraph_format.space_after = Pt(14)

doc.save(OUT)
print("已生成:", OUT)
for n in (1, 2, 3):
    cap, rows, foot = grab_table(n)
    print(f"  Table {n}: {len(rows)} 行 × {len(rows[0])} 列 | 表注 {len(foot.split())} 词")
